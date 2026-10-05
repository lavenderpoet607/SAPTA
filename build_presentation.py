import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

os.makedirs('mockup_assets', exist_ok=True)

FONT_REG = 'C:/Windows/Fonts/segoeui.ttf'
FONT_BOLD = 'C:/Windows/Fonts/segoeuib.ttf'
FONT_SEMI = 'C:/Windows/Fonts/seguisb.ttf'
FONT_MONO = 'C:/Windows/Fonts/consola.ttf'

def get_font(path, size):
    try:
        return ImageFont.truetype(path, size)
    except Exception:
        return ImageFont.load_default()

def create_phone_mockup(screen_type='dashboard', output_path='mockup_assets/mockup.png'):
    W, H = 540, 1100
    BEZEL = 16
    CORNER = 56
    SCREEN_CORNER = 42
    
    img = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    
    for i in range(12, 0, -1):
        alpha = int(28 * (1.0 - i / 12.0))
        draw.rounded_rectangle(
            [16 - i, 16 - i // 2, W - 16 + i, H - 16 + i],
            radius=CORNER + i,
            fill=(0, 0, 0, alpha)
        )
    
    draw.rounded_rectangle([16, 16, W - 16, H - 16], radius=CORNER, fill=(30, 41, 59, 255), outline=(51, 65, 85, 255), width=3)
    draw.rounded_rectangle([20, 20, W - 20, H - 20], radius=CORNER - 4, fill=(15, 23, 42, 255))
    
    screen_rect = [20 + BEZEL, 20 + BEZEL, W - 20 - BEZEL, H - 20 - BEZEL]
    s_x1, s_y1, s_x2, s_y2 = screen_rect
    s_w = s_x2 - s_x1
    s_h = s_y2 - s_y1
    
    screen = Image.new('RGBA', (s_w, s_h), (248, 250, 252, 255))
    s_draw = ImageDraw.Draw(screen)
    
    f_h1 = get_font(FONT_BOLD, 26)
    f_h2 = get_font(FONT_BOLD, 20)
    f_body = get_font(FONT_REG, 16)
    f_body_b = get_font(FONT_BOLD, 16)
    f_sub = get_font(FONT_REG, 13)
    f_sub_b = get_font(FONT_BOLD, 13)
    f_mono = get_font(FONT_MONO, 12)
    
    NAVY = (23, 20, 70, 255)
    BLUE_P = (30, 58, 138, 255)
    GREEN = (5, 150, 105, 255)
    ORANGE = (234, 88, 12, 255)
    AMBER = (217, 119, 6, 255)
    WHITE = (255, 255, 255, 255)
    GRAY_TEXT = (100, 116, 139, 255)
    DARK_TEXT = (15, 23, 42, 255)
    
    if screen_type == 'login':
        s_draw.rectangle([0, 0, s_w, 80], fill=NAVY)
        s_draw.text((s_w // 2, 50), "Masuk Akun", fill=WHITE, font=f_h2, anchor='mm')
        
        center_y = 200
        s_draw.ellipse([s_w // 2 - 45, center_y - 45, s_w // 2 + 45, center_y + 45], fill=(30, 58, 138, 30), outline=BLUE_P, width=2)
        s_draw.rectangle([s_w // 2 - 18, center_y - 20, s_w // 2 + 18, center_y + 20], fill=BLUE_P)
        s_draw.polygon([(s_w // 2 - 10, center_y), (s_w // 2 - 2, center_y + 8), (s_w // 2 + 12, center_y - 8)], fill=WHITE)
        
        s_draw.text((s_w // 2, center_y + 70), "SAPTA", fill=DARK_TEXT, font=get_font(FONT_BOLD, 30), anchor='mm')
        s_draw.text((s_w // 2, center_y + 100), "Sistem Absensi & Presensi Terpadu", fill=GRAY_TEXT, font=f_sub, anchor='mm')
        
        fy = center_y + 160
        s_draw.text((36, fy - 22), "Email", fill=DARK_TEXT, font=f_sub_b)
        s_draw.rounded_rectangle([32, fy, s_w - 32, fy + 52], radius=12, fill=WHITE, outline=(203, 213, 225, 255), width=2)
        s_draw.text((50, fy + 26), "peserta.ppkd@email.com", fill=DARK_TEXT, font=f_body, anchor='lm')
        
        fy += 84
        s_draw.text((36, fy - 22), "Kata Sandi", fill=DARK_TEXT, font=f_sub_b)
        s_draw.rounded_rectangle([32, fy, s_w - 32, fy + 52], radius=12, fill=WHITE, outline=(203, 213, 225, 255), width=2)
        s_draw.text((50, fy + 26), "••••••••••••", fill=DARK_TEXT, font=f_body, anchor='lm')
        
        fy += 90
        s_draw.rounded_rectangle([32, fy, s_w - 32, fy + 56], radius=12, fill=NAVY)
        s_draw.text((s_w // 2, fy + 28), "Masuk ke Akun", fill=WHITE, font=f_body_b, anchor='mm')
        
        s_draw.text((s_w // 2, fy + 100), "Belum punya akun? Daftar Sekarang", fill=BLUE_P, font=f_sub_b, anchor='mm')
        
    elif screen_type == 'dashboard':
        s_draw.rectangle([0, 0, s_w, 80], fill=NAVY)
        s_draw.text((32, 50), "SAPTA", fill=WHITE, font=f_h2, anchor='lm')
        s_draw.ellipse([s_w - 55, 35, s_w - 25, 65], fill=(255, 255, 255, 30))
        s_draw.text((s_w - 40, 50), "🌙", font=get_font(FONT_REG, 16), anchor='mm')
        
        for y in range(80, 240):
            ratio = (y - 80) / 160.0
            r = int(30 + ratio * (37 - 30))
            g = int(58 + ratio * (99 - 58))
            b = int(138 + ratio * (235 - 138))
            s_draw.line([(0, y), (s_w, y)], fill=(r, g, b, 255))
            
        s_draw.ellipse([32, 105, 82, 155], fill=(255, 255, 255, 40), outline=WHITE, width=2)
        s_draw.text((57, 130), "👤", font=get_font(FONT_REG, 24), anchor='mm')
        s_draw.text((95, 118), "Selamat Datang,", fill=(255, 255, 255, 200), font=f_sub)
        s_draw.text((95, 142), "Peserta PPKD Jakarta", fill=WHITE, font=f_h2)
        
        s_draw.rounded_rectangle([32, 180, s_w - 32, 220], radius=8, fill=(255, 255, 255, 35))
        s_draw.text((44, 200), "📅  Senin, 05 Oktober 2026", fill=WHITE, font=f_sub_b, anchor='lm')
        
        card_y = 255
        s_draw.rounded_rectangle([24, card_y, s_w - 24, card_y + 115], radius=16, fill=WHITE, outline=(226, 232, 240, 255), width=1)
        s_draw.text((40, card_y + 24), "📍 Lokasi Presensi Anda", fill=NAVY, font=f_sub_b, anchor='lm')
        s_draw.text((s_w - 44, card_y + 24), "🔄", font=get_font(FONT_REG, 14), anchor='rm')
        s_draw.text((40, card_y + 55), "Jl. Pemuda No. 45, Rawamangun, Jakarta Timur", fill=DARK_TEXT, font=f_sub, anchor='lm')
        s_draw.rounded_rectangle([40, card_y + 74, s_w - 40, card_y + 98], radius=6, fill=(241, 245, 249, 255))
        s_draw.text((50, card_y + 86), "GPS: -6.1934, 106.8821  •  Akurasi: ±3m", fill=GRAY_TEXT, font=f_mono, anchor='lm')
        
        btn_y = card_y + 130
        bw = (s_w - 48 - 14) // 2
        s_draw.rounded_rectangle([24, btn_y, 24 + bw, btn_y + 76], radius=14, fill=GREEN)
        s_draw.text((24 + bw // 2, btn_y + 26), "ABSEN MASUK", fill=WHITE, font=f_body_b, anchor='mm')
        s_draw.text((24 + bw // 2, btn_y + 50), "07:30 - 08:00 WIB", fill=(255, 255, 255, 210), font=f_sub, anchor='mm')
        
        s_draw.rounded_rectangle([24 + bw + 14, btn_y, s_w - 24, btn_y + 76], radius=14, fill=ORANGE)
        s_draw.text((24 + bw + 14 + bw // 2, btn_y + 26), "ABSEN PULANG", fill=WHITE, font=f_body_b, anchor='mm')
        s_draw.text((24 + bw + 14 + bw // 2, btn_y + 50), "16:00 WIB", fill=(255, 255, 255, 210), font=f_sub, anchor='mm')
        
        iz_y = btn_y + 88
        s_draw.rounded_rectangle([24, iz_y, s_w - 24, iz_y + 46], radius=10, fill=WHITE, outline=AMBER, width=2)
        s_draw.text((s_w // 2, iz_y + 23), "📝  Ajukan Izin / Keterangan Sakit", fill=AMBER, font=f_sub_b, anchor='mm')
        
        st_y = iz_y + 60
        s_draw.text((28, st_y), "Statistik Bulan Ini", fill=DARK_TEXT, font=f_sub_b)
        st_box_y = st_y + 16
        cw = (s_w - 48 - 20) // 3
        
        stats = [("18", "Masuk", GREEN, (209, 250, 229, 255)),
                 ("2", "Izin", AMBER, (254, 243, 199, 255)),
                 ("16", "Selesai", BLUE_P, (219, 234, 254, 255))]
        for i, (val, lbl, col, bgcol) in enumerate(stats):
            cx = 24 + i * (cw + 10)
            s_draw.rounded_rectangle([cx, st_box_y, cx + cw, st_box_y + 75], radius=12, fill=bgcol)
            s_draw.text((cx + cw // 2, st_box_y + 26), val, fill=col, font=f_h1, anchor='mm')
            s_draw.text((cx + cw // 2, st_box_y + 54), lbl, fill=col, font=f_sub_b, anchor='mm')
            
        nav_y = s_h - 75
        s_draw.rectangle([0, nav_y, s_w, s_h], fill=WHITE, outline=(226, 232, 240, 255), width=1)
        tabs = [("🏠", "Dashboard", True), ("📋", "Riwayat", False), ("👤", "Profil", False)]
        tw = s_w // 3
        for i, (ic, title, active) in enumerate(tabs):
            tx = i * tw + tw // 2
            col = NAVY if active else GRAY_TEXT
            s_draw.text((tx, nav_y + 25), ic, font=get_font(FONT_REG, 18), anchor='mm')
            s_draw.text((tx, nav_y + 52), title, fill=col, font=f_sub_b if active else f_sub, anchor='mm')
            if active:
                s_draw.line([(tx - 16, nav_y + 64), (tx + 16, nav_y + 64)], fill=NAVY, width=3)
                
    elif screen_type == 'map':
        s_draw.rectangle([0, 0, s_w, 80], fill=NAVY)
        s_draw.text((32, 50), "←  Detail Lokasi Presensi", fill=WHITE, font=f_h2, anchor='lm')
        
        map_h = 580
        s_draw.rectangle([0, 80, s_w, 80 + map_h], fill=(227, 238, 230, 255))
        for r_y in [180, 280, 420, 520]:
            s_draw.line([(0, r_y), (s_w, r_y)], fill=(255, 255, 255, 255), width=12)
        for r_x in [120, 280, 400]:
            s_draw.line([(r_x, 80), (r_x, 80 + map_h)], fill=(255, 255, 255, 255), width=12)
            
        s_draw.rounded_rectangle([140, 200, 260, 260], radius=8, fill=(240, 240, 240, 255))
        s_draw.text((200, 230), "Kantor PPKD", fill=GRAY_TEXT, font=f_sub, anchor='mm')
        
        cx, cy = s_w // 2, 340
        s_draw.ellipse([cx - 100, cy - 100, cx + 100, cy + 100], fill=(37, 99, 235, 40), outline=(37, 99, 235, 140), width=2)
        s_draw.ellipse([cx - 40, cy - 40, cx + 40, cy + 40], fill=(37, 99, 235, 70))
        s_draw.ellipse([cx - 16, cy - 16, cx + 16, cy + 16], fill=BLUE_P, outline=WHITE, width=3)
        s_draw.ellipse([cx - 6, cy - 6, cx + 6, cy + 6], fill=WHITE)
        
        sheet_y = 80 + map_h - 40
        s_draw.rounded_rectangle([0, sheet_y, s_w, s_h], radius=28, fill=WHITE, outline=(203, 213, 225, 255), width=1)
        s_draw.rounded_rectangle([s_w // 2 - 25, sheet_y + 12, s_w // 2 + 25, sheet_y + 16], radius=2, fill=(203, 213, 225, 255))
        
        sy = sheet_y + 40
        s_draw.text((28, sy), "Lokasi Terverifikasi", fill=GREEN, font=f_sub_b)
        s_draw.text((28, sy + 28), "Pusat Pelatihan Kerja Daerah", fill=DARK_TEXT, font=f_h2)
        s_draw.text((28, sy + 58), "Jl. Pemuda No. 45, Rawamangun, Kec. Pulo Gadung", fill=GRAY_TEXT, font=f_sub)
        
        info_y = sy + 90
        s_draw.rounded_rectangle([24, info_y, s_w - 24, info_y + 80], radius=12, fill=(248, 250, 252, 255), outline=(226, 232, 240, 255), width=1)
        s_draw.text((40, info_y + 24), "Koordinat: -6.193421, 106.882190", fill=DARK_TEXT, font=f_mono)
        s_draw.text((40, info_y + 50), "Radius Presensi: Valid (Dalam 50 meter)", fill=GREEN, font=f_sub_b)
        
        btn_y = info_y + 95
        s_draw.rounded_rectangle([24, btn_y, s_w - 24, btn_y + 52], radius=12, fill=NAVY)
        s_draw.text((s_w // 2, btn_y + 26), "Konfirmasi Kehadiran Disini", fill=WHITE, font=f_body_b, anchor='mm')
        
    elif screen_type == 'history':
        s_draw.rectangle([0, 0, s_w, 80], fill=NAVY)
        s_draw.text((32, 50), "Riwayat Absensi", fill=WHITE, font=f_h2, anchor='lm')
        s_draw.text((s_w - 32, 50), "🔍", font=get_font(FONT_REG, 18), anchor='rm')
        
        chip_y = 100
        chips = [("Semua (20)", True), ("Masuk (18)", False), ("Izin (2)", False)]
        cx = 24
        for label, is_act in chips:
            cw = 110 if "Semua" in label else 95
            bg = NAVY if is_act else WHITE
            fg = WHITE if is_act else GRAY_TEXT
            s_draw.rounded_rectangle([cx, chip_y, cx + cw, chip_y + 36], radius=18, fill=bg, outline=NAVY if is_act else (203, 213, 225, 255), width=1)
            s_draw.text((cx + cw // 2, chip_y + 18), label, fill=fg, font=f_sub_b, anchor='mm')
            cx += cw + 10
            
        list_y = 155
        items = [
            ("Senin, 05 Okt 2026", "07:48 WIB", "Masuk Tepat Waktu", GREEN, "Jl. Pemuda No. 45, Rawamangun"),
            ("Jumat, 02 Okt 2026", "16:05 WIB", "Pulang Selesai", ORANGE, "Jl. Pemuda No. 45, Rawamangun"),
            ("Kamis, 01 Okt 2026", "08:15 WIB", "Izin Sakit", AMBER, "Surat Keterangan Dokter Terlampir"),
            ("Rabu, 30 Sep 2026", "07:42 WIB", "Masuk Tepat Waktu", GREEN, "Jl. Pemuda No. 45, Rawamangun"),
            ("Selasa, 29 Sep 2026", "07:50 WIB", "Masuk Tepat Waktu", GREEN, "Jl. Pemuda No. 45, Rawamangun"),
        ]
        
        for date, time, status, col, addr in items:
            s_draw.rounded_rectangle([24, list_y, s_w - 24, list_y + 100], radius=14, fill=WHITE, outline=(226, 232, 240, 255), width=1)
            s_draw.text((40, list_y + 22), date, fill=DARK_TEXT, font=f_body_b)
            
            bw = 140
            s_draw.rounded_rectangle([s_w - 38 - bw, list_y + 12, s_w - 38, list_y + 36], radius=12, fill=col)
            s_draw.text((s_w - 38 - bw // 2, list_y + 24), status, fill=WHITE, font=get_font(FONT_BOLD, 11), anchor='mm')
            
            s_draw.text((40, list_y + 52), f"⏰ Waktu: {time}", fill=GRAY_TEXT, font=f_sub)
            s_draw.text((40, list_y + 76), f"📍 {addr}", fill=GRAY_TEXT, font=f_sub)
            list_y += 114
            
        nav_y = s_h - 75
        s_draw.rectangle([0, nav_y, s_w, s_h], fill=WHITE, outline=(226, 232, 240, 255), width=1)
        tabs = [("🏠", "Dashboard", False), ("📋", "Riwayat", True), ("👤", "Profil", False)]
        tw = s_w // 3
        for i, (ic, title, active) in enumerate(tabs):
            tx = i * tw + tw // 2
            col = NAVY if active else GRAY_TEXT
            s_draw.text((tx, nav_y + 25), ic, font=get_font(FONT_REG, 18), anchor='mm')
            s_draw.text((tx, nav_y + 52), title, fill=col, font=f_sub_b if active else f_sub, anchor='mm')
            if active:
                s_draw.line([(tx - 16, nav_y + 64), (tx + 16, nav_y + 64)], fill=NAVY, width=3)
                
    elif screen_type == 'dark':
        s_draw.rectangle([0, 0, s_w, s_h], fill=(15, 23, 42, 255))
        s_draw.rectangle([0, 0, s_w, 80], fill=(23, 20, 70, 255))
        s_draw.text((32, 50), "SAPTA (Dark Mode)", fill=WHITE, font=f_h2, anchor='lm')
        s_draw.text((s_w - 40, 50), "☀️", font=get_font(FONT_REG, 16), anchor='mm')
        
        card_y = 100
        s_draw.rounded_rectangle([24, card_y, s_w - 24, card_y + 110], radius=16, fill=(30, 41, 59, 255), outline=(51, 65, 85, 255), width=1)
        s_draw.ellipse([44, card_y + 25, 104, card_y + 85], fill=(51, 65, 85, 255))
        s_draw.text((74, card_y + 55), "👤", font=get_font(FONT_REG, 26), anchor='mm')
        s_draw.text((120, card_y + 40), "Peserta PPKD Jakarta", fill=WHITE, font=f_h2)
        s_draw.text((120, card_y + 68), "Web & Mobile App • Batch 1", fill=(148, 163, 184, 255), font=f_sub)
        
        iz_y = card_y + 130
        s_draw.rounded_rectangle([24, iz_y, s_w - 24, iz_y + 360], radius=16, fill=(30, 41, 59, 255), outline=(51, 65, 85, 255), width=1)
        s_draw.text((40, iz_y + 30), "Formulir Pengajuan Izin", fill=WHITE, font=f_h2)
        s_draw.text((40, iz_y + 58), "Isi data keterangan kehadiran Anda", fill=(148, 163, 184, 255), font=f_sub)
        
        fy = iz_y + 90
        s_draw.text((40, fy), "Kategori Izin", fill=WHITE, font=f_sub_b)
        s_draw.rounded_rectangle([40, fy + 16, s_w - 40, fy + 60], radius=10, fill=(15, 23, 42, 255), outline=(51, 65, 85, 255), width=1)
        s_draw.text((54, fy + 38), "Izin Sakit (Disertai Surat Dokter)", fill=WHITE, font=f_body, anchor='lm')
        
        fy += 80
        s_draw.text((40, fy), "Alasan / Keterangan", fill=WHITE, font=f_sub_b)
        s_draw.rounded_rectangle([40, fy + 16, s_w - 40, fy + 110], radius=10, fill=(15, 23, 42, 255), outline=(51, 65, 85, 255), width=1)
        s_draw.text((54, fy + 44), "Mengalami demam tinggi dan flu,", fill=(203, 213, 225, 255), font=f_body)
        s_draw.text((54, fy + 70), "disarankan istirahat oleh dokter.", fill=(203, 213, 225, 255), font=f_body)
        
        btn_y = fy + 126
        s_draw.rounded_rectangle([40, btn_y, s_w - 40, btn_y + 52], radius=12, fill=AMBER)
        s_draw.text((s_w // 2, btn_y + 26), "Kirim Surat Izin", fill=WHITE, font=f_body_b, anchor='mm')
        
        nav_y = s_h - 75
        s_draw.rectangle([0, nav_y, s_w, s_h], fill=(23, 20, 70, 255), outline=(51, 65, 85, 255), width=1)
        tabs = [("🏠", "Dashboard", False), ("📋", "Riwayat", False), ("👤", "Profil", True)]
        tw = s_w // 3
        for i, (ic, title, active) in enumerate(tabs):
            tx = i * tw + tw // 2
            col = WHITE if active else (148, 163, 184, 255)
            s_draw.text((tx, nav_y + 25), ic, font=get_font(FONT_REG, 18), anchor='mm')
            s_draw.text((tx, nav_y + 52), title, fill=col, font=f_sub_b if active else f_sub, anchor='mm')
            if active:
                s_draw.line([(tx - 16, nav_y + 64), (tx + 16, nav_y + 64)], fill=WHITE, width=3)
                
    mask = Image.new('L', (s_w, s_h), 0)
    mask_draw = ImageDraw.Draw(mask)
    mask_draw.rounded_rectangle([0, 0, s_w, s_h], radius=SCREEN_CORNER, fill=255)
    
    img.paste(screen, (s_x1, s_y1), mask)
    
    island_w, island_h = 130, 32
    i_x1 = W // 2 - island_w // 2
    i_y1 = s_y1 + 10
    draw.rounded_rectangle([i_x1, i_y1, i_x1 + island_w, i_y1 + island_h], radius=16, fill=(0, 0, 0, 255))
    draw.ellipse([i_x1 + island_w - 28, i_y1 + 8, i_x1 + island_w - 12, i_y1 + 24], fill=(15, 23, 42, 255))
    draw.ellipse([i_x1 + island_w - 22, i_y1 + 14, i_x1 + island_w - 18, i_y1 + 18], fill=(30, 41, 59, 255))
    
    h_bar_w, h_bar_h = 140, 5
    hb_x1 = W // 2 - h_bar_w // 2
    hb_y1 = s_y2 - 14
    draw.rounded_rectangle([hb_x1, hb_y1, hb_x1 + h_bar_w, hb_y1 + h_bar_h], radius=3, fill=(148, 163, 184, 180))
    
    img.save(output_path, 'PNG')
    print(f"Generated mockup: {output_path}")

def build_presentation():
    create_phone_mockup('login', 'mockup_assets/mockup_login.png')
    create_phone_mockup('dashboard', 'mockup_assets/mockup_dashboard.png')
    create_phone_mockup('map', 'mockup_assets/mockup_map.png')
    create_phone_mockup('history', 'mockup_assets/mockup_history.png')
    create_phone_mockup('dark', 'mockup_assets/mockup_dark.png')
    
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    C_MIDNIGHT = RGBColor(23, 20, 70)
    C_NAVY = RGBColor(30, 58, 138)
    C_ROYAL = RGBColor(37, 99, 235)
    C_EMERALD = RGBColor(5, 150, 105)
    C_ORANGE = RGBColor(234, 88, 12)
    C_AMBER = RGBColor(217, 119, 6)
    C_BG_LIGHT = RGBColor(248, 250, 252)
    C_DARK_SLATE = RGBColor(15, 23, 42)
    C_WHITE = RGBColor(255, 255, 255)
    C_GRAY = RGBColor(100, 116, 139)
    C_CARD_BORDER = RGBColor(226, 232, 240)
    
    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.color.rgb = color
        return bg

    # -------------------------------------------------------------
    # SLIDE 1: COVER
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, C_MIDNIGHT)
    
    # Accent shape
    acc = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(3.8), Inches(0.42))
    acc.fill.solid()
    acc.fill.fore_color.rgb = C_ROYAL
    acc.line.fill.background()
    acc.text_frame.text = "APLIKASI MOBILE PRESENSI DIGITAL"
    acc.text_frame.paragraphs[0].font.size = Pt(11)
    acc.text_frame.paragraphs[0].font.bold = True
    acc.text_frame.paragraphs[0].font.color.rgb = C_WHITE
    acc.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    t1_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(7.2), Inches(4.5))
    tf1 = t1_box.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "SAPTA"
    p.font.size = Pt(58)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    
    p2 = tf1.add_paragraph()
    p2.text = "Sistem Absensi & Presensi Terpadu"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(96, 165, 250)
    p2.space_before = Pt(8)
    
    p3 = tf1.add_paragraph()
    p3.text = "Solusi pencatatan kehadiran modern berbasis mobile Flutter dengan validasi titik koordinat GPS real-time, Google Maps terintegrasi, dan rekapitulasi kehadiran instan."
    p3.font.size = Pt(15)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(20)
    
    p4 = tf1.add_paragraph()
    p4.text = "🏛 Pusat Pelatihan Kerja Daerah (PPKD)  •  Tahun 2026"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = C_WHITE
    p4.space_before = Pt(36)
    
    s1.shapes.add_picture('mockup_assets/mockup_login.png', Inches(8.6), Inches(0.6), height=Inches(6.3))
    
    # -------------------------------------------------------------
    # SLIDE 2: LATAR BELAKANG & MASALAH
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, C_BG_LIGHT)
    
    # Header tag
    t2_tag = s2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.0), Inches(0.4))
    t2_tag.text_frame.text = "LATAR BELAKANG & MASALAH"
    t2_tag.text_frame.paragraphs[0].font.size = Pt(12)
    t2_tag.text_frame.paragraphs[0].font.bold = True
    t2_tag.text_frame.paragraphs[0].font.color.rgb = C_NAVY
    
    t2_title = s2.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(7.5), Inches(0.8))
    t2_title.text_frame.text = "Tantangan Sistem Presensi Manual"
    t2_title.text_frame.paragraphs[0].font.size = Pt(28)
    t2_title.text_frame.paragraphs[0].font.bold = True
    t2_title.text_frame.paragraphs[0].font.color.rgb = C_DARK_SLATE
    
    problems = [
        ("Rawan Manipulasi & Titip Absen", "Pencatatan tanda tangan manual atau mesin fingerprint statis rentan manipulasi dan tidak memverifikasi keberadaan peserta secara aktual di lokasi pelatihan."),
        ("Verifikasi Lokasi yang Lemah", "Sulit mengetahui apakah peserta pelatihan atau karyawan benar-benar berada di lingkungan instansi saat waktu absensi masuk maupun pulang."),
        ("Rekapitulasi Lambat & Tidak Efisien", "Proses pengumpulan data absensi secara konvensional memakan waktu berhari-hari dan rawan kesalahan kalkulasi data kehadiran.")
    ]
    
    card_y = 1.8
    for title, desc in problems:
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(card_y), Inches(7.2), Inches(1.2))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = "⚠️  " + title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_ORANGE
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_GRAY
        p2.space_before = Pt(4)
        
        card_y += 1.35
        
    sol_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.95), Inches(7.2), Inches(0.95))
    sol_card.fill.solid()
    sol_card.fill.fore_color.rgb = C_MIDNIGHT
    sol_card.line.fill.background()
    stf = sol_card.text_frame
    stf.word_wrap = True
    stf.vertical_anchor = MSO_ANCHOR.MIDDLE
    sp = stf.paragraphs[0]
    sp.text = "💡 Solusi SAPTA: Otomasi presensi mobile dengan validasi GPS instan & integrasi cloud real-time."
    sp.font.size = Pt(13)
    sp.font.bold = True
    sp.font.color.rgb = C_WHITE
    sp.alignment = PP_ALIGN.CENTER
    
    s2.shapes.add_picture('mockup_assets/mockup_login.png', Inches(8.8), Inches(0.6), height=Inches(6.3))

    # -------------------------------------------------------------
    # SLIDE 3: OVERVIEW APLIKASI
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, C_BG_LIGHT)
    
    t3_tag = s3.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.0), Inches(0.4))
    t3_tag.text_frame.text = "PENGENALAN APLIKASI"
    t3_tag.text_frame.paragraphs[0].font.size = Pt(12)
    t3_tag.text_frame.paragraphs[0].font.bold = True
    t3_tag.text_frame.paragraphs[0].font.color.rgb = C_NAVY
    
    t3_title = s3.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(7.5), Inches(0.8))
    t3_title.text_frame.text = "SAPTA: Presensi Digital Modern"
    t3_title.text_frame.paragraphs[0].font.size = Pt(28)
    t3_title.text_frame.paragraphs[0].font.bold = True
    t3_title.text_frame.paragraphs[0].font.color.rgb = C_DARK_SLATE
    
    overview_items = [
        ("Dashboard Informatif", "Menampilkan ucapan selamat datang personal, tanggal hari ini, ringkasan jam presensi, dan status kehadiran langsung di beranda utama."),
        ("Pencatatan Presensi Sekali Sentuh", "Tombol cepat Absen Masuk dan Absen Pulang dengan warna kontras yang mempermudah interaksi pengguna dalam hitungan detik."),
        ("Statistik Kehadiran Akumulatif", "Kartu counter ringkas yang memantau total kehadiran (Masuk, Izin, Selesai) untuk transparansi evaluasi kedisiplinan peserta."),
        ("Navigasi Bersih & Intuitif", "Bottom Navigation Bar 3 tab (Dashboard, Riwayat, Profil) yang memudahkan eksplorasi seluruh modul aplikasi.")
    ]
    
    card_y = 1.8
    for title, desc in overview_items:
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(card_y), Inches(7.2), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = "✓  " + title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_NAVY
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = C_GRAY
        p2.space_before = Pt(3)
        
        card_y += 1.2
        
    s3.shapes.add_picture('mockup_assets/mockup_dashboard.png', Inches(8.8), Inches(0.6), height=Inches(6.3))

    # -------------------------------------------------------------
    # SLIDE 4: GPS & GOOGLE MAPS
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, C_BG_LIGHT)
    
    t4_tag = s4.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.0), Inches(0.4))
    t4_tag.text_frame.text = "FITUR UNGGULAN 1"
    t4_tag.text_frame.paragraphs[0].font.size = Pt(12)
    t4_tag.text_frame.paragraphs[0].font.bold = True
    t4_tag.text_frame.paragraphs[0].font.color.rgb = C_ROYAL
    
    t4_title = s4.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(7.5), Inches(0.8))
    t4_title.text_frame.text = "Presensi GPS & Integrasi Peta"
    t4_title.text_frame.paragraphs[0].font.size = Pt(28)
    t4_title.text_frame.paragraphs[0].font.bold = True
    t4_title.text_frame.paragraphs[0].font.color.rgb = C_DARK_SLATE
    
    gps_items = [
        ("Deteksi Koordinat Presisi Tinggi", "Membaca Latitude dan Longitude secara instan dari satelit GPS dengan akurasi radius tinggi saat tombol ditekan."),
        ("Reverse Geocoding Otomatis", "Titik koordinat langsung diterjemahkan menjadi nama jalan, kelurahan, dan kota yang mudah dibaca oleh admin dan peserta."),
        ("Visualisasi Peta Interaktif", "Didukung Google Maps dengan penanda pin lokasi dinamis dan lingkaran geofencing untuk memastikan kehadiran di area yang ditentukan."),
        ("Dialog Verifikasi Lokasi", "Pengguna mendapatkan konfirmasi alamat dan koordinat sebelum presensi final dikirimkan ke server.")
    ]
    
    card_y = 1.8
    for title, desc in gps_items:
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(card_y), Inches(7.2), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = "📍  " + title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_ROYAL
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = C_GRAY
        p2.space_before = Pt(3)
        
        card_y += 1.2
        
    s4.shapes.add_picture('mockup_assets/mockup_map.png', Inches(8.8), Inches(0.6), height=Inches(6.3))

    # -------------------------------------------------------------
    # SLIDE 5: RIWAYAT & MONITORING
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, C_BG_LIGHT)
    
    t5_tag = s5.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.0), Inches(0.4))
    t5_tag.text_frame.text = "FITUR UNGGULAN 2"
    t5_tag.text_frame.paragraphs[0].font.size = Pt(12)
    t5_tag.text_frame.paragraphs[0].font.bold = True
    t5_tag.text_frame.paragraphs[0].font.color.rgb = C_EMERALD
    
    t5_title = s5.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(7.5), Inches(0.8))
    t5_title.text_frame.text = "Riwayat & Monitoring Kehadiran"
    t5_title.text_frame.paragraphs[0].font.size = Pt(28)
    t5_title.text_frame.paragraphs[0].font.bold = True
    t5_title.text_frame.paragraphs[0].font.color.rgb = C_DARK_SLATE
    
    hist_items = [
        ("Log Kehadiran Kronologis", "Daftar histori presensi tersusun rapi dari tanggal terbaru lengkap dengan informasi jam masuk, jam pulang, dan alamat lokasi."),
        ("Status Berkode Warna", "Badge status visual yang sangat jelas: Hijau (Masuk Tepat Waktu), Oranye (Pulang Selesai), dan Kuning (Izin Sakit / Keperluan)."),
        ("Filter Kategori Praktis", "Filter instan untuk menampilkan kategori 'Semua', 'Masuk', atau 'Izin' hanya dengan satu ketukan."),
        ("Detail Lokasi Per Riwayat", "Setiap catatan riwayat dapat dibuka kembali untuk melihat rincian koordinat peta saat presensi dilakukan.")
    ]
    
    card_y = 1.8
    for title, desc in hist_items:
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(card_y), Inches(7.2), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = "📋  " + title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_EMERALD
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = C_GRAY
        p2.space_before = Pt(3)
        
        card_y += 1.2
        
    s5.shapes.add_picture('mockup_assets/mockup_history.png', Inches(8.8), Inches(0.6), height=Inches(6.3))

    # -------------------------------------------------------------
    # SLIDE 6: FORM IZIN & DARK MODE
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, C_BG_LIGHT)
    
    t6_tag = s6.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.0), Inches(0.4))
    t6_tag.text_frame.text = "FITUR UNGGULAN 3"
    t6_tag.text_frame.paragraphs[0].font.size = Pt(12)
    t6_tag.text_frame.paragraphs[0].font.bold = True
    t6_tag.text_frame.paragraphs[0].font.color.rgb = C_AMBER
    
    t6_title = s6.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(7.5), Inches(0.8))
    t6_title.text_frame.text = "Form Izin & Tampilan Dark Mode"
    t6_title.text_frame.paragraphs[0].font.size = Pt(28)
    t6_title.text_frame.paragraphs[0].font.bold = True
    t6_title.text_frame.paragraphs[0].font.color.rgb = C_DARK_SLATE
    
    dark_items = [
        ("Formulir Izin Cepat & Terstruktur", "Pengguna dapat mengajukan izin sakit atau keperluan darurat disertai keterangan alasan tertulis secara langsung dari HP."),
        ("Dukungan Tema Gelap (Dark Mode)", "Tersedia switch toggle dark mode instan yang ramah di mata saat malam hari dan menghemat konsumsi daya baterai."),
        ("Kontras Adaptif Sempurna", "Seluruh teks, container alamat, dan tombol tetap kontras dan nyaman dibaca baik pada mode terang maupun mode gelap."),
        ("Manajemen Profil Pengguna", "Halaman profil lengkap untuk melihat informasi nama, kejuruan pelatihan, ubah preferensi, hingga opsi logout yang aman.")
    ]
    
    card_y = 1.8
    for title, desc in dark_items:
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(card_y), Inches(7.2), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = "🌙  " + title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_AMBER
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = C_GRAY
        p2.space_before = Pt(3)
        
        card_y += 1.2
        
    s6.shapes.add_picture('mockup_assets/mockup_dark.png', Inches(8.8), Inches(0.6), height=Inches(6.3))

    # -------------------------------------------------------------
    # SLIDE 7: KEUNGGULAN SISTEM
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, C_BG_LIGHT)
    
    t7_tag = s7.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.0), Inches(0.4))
    t7_tag.text_frame.text = "NILAI TAMBAH"
    t7_tag.text_frame.paragraphs[0].font.size = Pt(12)
    t7_tag.text_frame.paragraphs[0].font.bold = True
    t7_tag.text_frame.paragraphs[0].font.color.rgb = C_NAVY
    
    t7_title = s7.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(7.5), Inches(0.8))
    t7_title.text_frame.text = "Mengapa Memilih SAPTA?"
    t7_title.text_frame.paragraphs[0].font.size = Pt(28)
    t7_title.text_frame.paragraphs[0].font.bold = True
    t7_title.text_frame.paragraphs[0].font.color.rgb = C_DARK_SLATE
    
    advantages = [
        ("Akurasi Geofencing 100%", "Memastikan absensi hanya dapat dilakukan jika pengguna benar-benar berada di radius koordinat yang ditentukan."),
        ("Efisien & Bebas Antrean", "Menghilangkan antrean fisik di depan mesin fingerprint. Ratusan peserta dapat absen serentak dalam hitungan detik."),
        ("Desain UI/UX Sangat Ramah", "Antarmuka modern, bersih, intuitif, dan responsif yang mudah digunakan oleh siapa saja."),
        ("Data Terintegrasi & Real-Time", "Data kehadiran langsung terkirim ke server backend cloud, siap divalidasi dan dianalisis kapan saja.")
    ]
    
    card_y = 1.8
    for title, desc in advantages:
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(card_y), Inches(7.2), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = "⭐  " + title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_NAVY
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = C_GRAY
        p2.space_before = Pt(3)
        
        card_y += 1.2
        
    s7.shapes.add_picture('mockup_assets/mockup_dashboard.png', Inches(8.8), Inches(0.6), height=Inches(6.3))

    # -------------------------------------------------------------
    # SLIDE 8: PENUTUP & Q&A
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, C_MIDNIGHT)
    
    t8_box = s8.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(7.2), Inches(4.5))
    tf8 = t8_box.text_frame
    tf8.word_wrap = True
    
    p = tf8.paragraphs[0]
    p.text = "Terima Kasih"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    
    p2 = tf8.add_paragraph()
    p2.text = "\"Meningkatkan Kedisiplinan, Integritas, dan Efisiensi Bersama SAPTA.\""
    p2.font.size = Pt(20)
    p2.font.italic = True
    p2.font.color.rgb = RGBColor(147, 197, 253)
    p2.space_before = Pt(16)
    
    p3 = tf8.add_paragraph()
    p3.text = "Ada pertanyaan seputar alur sistem, integrasi GPS, atau implementasi aplikasi SAPTA?"
    p3.font.size = Pt(15)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(24)
    
    qa_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.8), Inches(6.0), Inches(0.8))
    qa_card.fill.solid()
    qa_card.fill.fore_color.rgb = C_ROYAL
    qa_card.line.fill.background()
    qtf = qa_card.text_frame
    qtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    qp = qtf.paragraphs[0]
    qp.text = "💬  SESI TANYA JAWAB (Q & A) DIBUKA"
    qp.font.size = Pt(14)
    qp.font.bold = True
    qp.font.color.rgb = C_WHITE
    qp.alignment = PP_ALIGN.CENTER
    
    s8.shapes.add_picture('mockup_assets/mockup_dashboard.png', Inches(8.6), Inches(0.6), height=Inches(6.3))
    
    output_pptx = 'SAPTA_Presentasi_Mockup.pptx'
    prs.save(output_pptx)
    print(f"Presentation saved successfully to: {os.path.abspath(output_pptx)}")

if __name__ == '__main__':
    build_presentation()
